# -*- coding: utf-8 -*-
"""
Genera el entregable en Word (.docx) del modulo
"Fundamentos de Bases de Datos Relacionales" - Alke Wallet.

Uso:  python generar_documento.py
Requiere: python-docx  (pip install python-docx)
Salida:  Alke_Wallet_Bases_de_Datos_Erick_Noguera.docx
"""
import os
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

AQUI = os.path.dirname(os.path.abspath(__file__))
MORADO = RGBColor(0x4F, 0x2D, 0x7F)

# --------------------------------------------------------------------- helpers
def shade(paragraph, fill):
    pPr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill)
    pPr.append(shd)

def border(paragraph, color="4F2D7F", size="6"):
    pPr = paragraph._p.get_or_add_pPr()
    bdr = OxmlElement('w:pBdr')
    for edge in ('top', 'left', 'bottom', 'right'):
        e = OxmlElement(f'w:{edge}')
        e.set(qn('w:val'), 'single')
        e.set(qn('w:sz'), size)
        e.set(qn('w:space'), '4')
        e.set(qn('w:color'), color)
        bdr.append(e)
    pPr.append(bdr)

def code_block(doc, texto):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.left_indent = Pt(6)
    shade(p, "F2F0F7")
    lineas = texto.strip("\n").split("\n")
    for i, ln in enumerate(lineas):
        run = p.add_run(ln)
        run.font.name = "Consolas"
        run.font.size = Pt(9)
        if i < len(lineas) - 1:
            run.add_break()
    return p

def capture_box(doc, texto):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(10)
    shade(p, "FFF7E6")
    border(p, "A15C00")
    r = p.add_run("[ CAPTURA DE PANTALLA ]  " + texto)
    r.bold = True
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor(0xA1, 0x5C, 0x00)
    return p

def result_table(doc, headers, rows, nota="Resultado esperado con los datos de prueba:"):
    if nota:
        np = doc.add_paragraph()
        rr = np.add_run(nota)
        rr.italic = True
        rr.font.size = Pt(9)
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Light Grid Accent 1"
    for j, h in enumerate(headers):
        c = t.rows[0].cells[j]
        c.text = str(h)
        for pr in c.paragraphs:
            for rn in pr.runs:
                rn.bold = True
                rn.font.size = Pt(9)
    for row in rows:
        cells = t.add_row().cells
        for j, val in enumerate(row):
            cells[j].text = str(val)
            for pr in cells[j].paragraphs:
                for rn in pr.runs:
                    rn.font.size = Pt(9)
    doc.add_paragraph()
    return t

def info_table(doc, headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Light Grid Accent 1"
    for j, h in enumerate(headers):
        c = t.rows[0].cells[j]
        c.text = str(h)
        for pr in c.paragraphs:
            for rn in pr.runs:
                rn.bold = True
                rn.font.size = Pt(9)
    for row in rows:
        cells = t.add_row().cells
        for j, val in enumerate(row):
            cells[j].text = str(val)
            for pr in cells[j].paragraphs:
                for rn in pr.runs:
                    rn.font.size = Pt(9)
    doc.add_paragraph()
    return t

def h(doc, texto, nivel=1):
    p = doc.add_heading(texto, level=nivel)
    for r in p.runs:
        r.font.color.rgb = MORADO
    return p

def para(doc, texto):
    p = doc.add_paragraph(texto)
    p.paragraph_format.space_after = Pt(6)
    return p

# --------------------------------------------------------------------- documento
doc = Document()
base = doc.styles["Normal"]
base.font.name = "Calibri"
base.font.size = Pt(11)

# ---- Portada
t = doc.add_paragraph()
t.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = t.add_run("ALKE WALLET")
r.bold = True
r.font.size = Pt(30)
r.font.color.rgb = MORADO

s = doc.add_paragraph()
s.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = s.add_run("Fundamentos de Bases de Datos Relacionales")
r.font.size = Pt(15)

for _ in range(3):
    doc.add_paragraph()

for linea in [
    "Entregable: diseno del modelo conceptual, relaciones y creacion de la base de datos",
    "Motor: MySQL 8  (ejecutado en sqliteonline.com - modo MySQL 8)",
    "Autor: Erick Noguera",
    "Repositorio: github.com/ErickNoguera/AlkeWallet_ErickNoguera  (carpeta /database)",
]:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(linea).font.size = Pt(11)

doc.add_page_break()

# ---- Indice
h(doc, "Contenido", 1)
for item in [
    "1. Introduccion y objetivo",
    "2. Leccion 1 - Las bases de datos relacionales",
    "3. Leccion 2 - Consultas a una o varias tablas",
    "4. Leccion 3 - Manipulacion de datos (DML) y transaccionalidad (ACID)",
    "5. Leccion 4 - Definicion de tablas (DDL)",
    "6. Leccion 5 - El modelo entidad-relacion",
    "7. Consultas obligatorias del enunciado (resumen)",
    "Anexo A - Script completo AlkeWallet.sql",
    "Anexo B - Checklist de capturas de pantalla",
]:
    doc.add_paragraph(item, style="List Bullet")
doc.add_page_break()

# =====================================================================
h(doc, "1. Introduccion y objetivo", 1)
para(doc, "El objetivo de esta evaluacion es disenar el modelo conceptual de la "
          "billetera digital Alke Wallet, definir las relaciones entre sus "
          "entidades y crear la base de datos relacional que almacena usuarios, "
          "monedas y transacciones, garantizando la coherencia y la integridad "
          "de los datos.")
para(doc, "La base de datos se llama AlkeWallet y contiene tres entidades: "
          "usuario, moneda y transaccion. Todas las sentencias SQL usadas para "
          "crear y manipular la base estan recogidas en este documento y en el "
          "archivo AlkeWallet.sql. Nota: el script .sql esta ordenado para "
          "ejecucion real (crear base -> DDL -> DML); este documento sigue el "
          "orden de lecciones del enunciado.")
para(doc, "Decision de diseno: el enunciado exige una consulta sobre 'la moneda "
          "elegida por un usuario' y una clave foranea de transaccion hacia la "
          "moneda, pero no lista el atributo currency_id en las entidades. Por "
          "ello se anade currency_id como clave foranea tanto en usuario (moneda "
          "del monedero) como en transaccion (moneda del movimiento). El atributo "
          "'contrasena' se nombra sin tilde por portabilidad del identificador.")

# =====================================================================
h(doc, "2. Leccion 1 - Las bases de datos relacionales", 1)

h(doc, "2.1 Que es una base de datos relacional y sus ventajas", 2)
para(doc, "Una base de datos relacional organiza la informacion en tablas "
          "(relaciones) formadas por filas (registros) y columnas (atributos). "
          "Cada tabla tiene una clave primaria que identifica de forma unica a "
          "cada fila, y puede vincularse con otras tablas mediante claves "
          "foraneas. El acceso y la manipulacion se realizan con el lenguaje SQL.")
for v in [
    "Integridad: las claves primarias, foraneas y las restricciones (NOT NULL, "
    "UNIQUE, CHECK) impiden datos incoherentes.",
    "No redundancia: la normalizacion evita repetir informacion y las anomalias "
    "de insercion, actualizacion y borrado.",
    "Consultas potentes y declarativas: SQL permite combinar tablas (JOIN), "
    "agrupar y agregar sin describir el 'como'.",
    "Transaccionalidad ACID: operaciones atomicas y duraderas, clave para un "
    "sistema financiero como una wallet.",
    "Estandar y portabilidad: el modelo relacional y SQL son un estandar "
    "soportado por multiples motores.",
]:
    doc.add_paragraph(v, style="List Bullet")

h(doc, "2.2 Tabla comparativa: RDBMS libres vs. comerciales", 2)
info_table(doc,
    ["Aspecto", "Libres (MySQL, PostgreSQL, MariaDB, SQLite)", "Comerciales (Oracle, SQL Server, Db2)"],
    [
        ["Licencia", "Codigo abierto (GPL / permisiva), sin costo", "Propietaria, de pago"],
        ["Costo", "Gratuito; se paga solo el soporte opcional", "Licencia por nucleo/usuario y soporte de alto costo"],
        ["Soporte", "Comunidad amplia; soporte comercial opcional", "Soporte oficial 24/7 incluido"],
        ["Herramientas", "Workbench, DBeaver, pgAdmin (terceros)", "Suites propias (SQL Developer, SSMS)"],
        ["Escalabilidad", "Alta con ajuste/replicacion/extensiones", "Muy alta, con herramientas empresariales integradas"],
        ["Casos de uso tipicos", "Web, startups, proyectos pequenos y medianos", "Banca, ERP, sistemas criticos de gran volumen"],
    ])

h(doc, "2.3 Creacion y verificacion de la base de datos", 2)
code_block(doc, """
CREATE DATABASE IF NOT EXISTS AlkeWallet
  CHARACTER SET utf8mb4
  COLLATE utf8mb4_0900_ai_ci;

USE AlkeWallet;

SHOW DATABASES;      -- verifica que AlkeWallet aparece en la lista
SHOW TABLES;         -- lista de tablas de la base activa
DESCRIBE usuario;    -- estructura de una tabla (tras crearla)
""")
capture_box(doc, "Pega la captura del primer CREATE DATABASE ejecutado con exito "
                 "en sqliteonline.com y del resultado de SHOW DATABASES;.")

# =====================================================================
h(doc, "3. Leccion 2 - Consultas a una o varias tablas", 1)
para(doc, "Las consultas siguientes se ejecutan sobre los cinco registros de "
          "prueba cargados en la Leccion 3.")

h(doc, "3.1 SELECT basicas sobre usuario", 2)
code_block(doc, """
SELECT * FROM usuario;
SELECT nombre, correo, saldo FROM usuario;
""")

h(doc, "3.2 Filtros con WHERE y operadores logicos", 2)
code_block(doc, """
SELECT nombre, saldo
FROM usuario
WHERE saldo >= 100000 AND currency_id = 1;

SELECT nombre, correo
FROM usuario
WHERE nombre LIKE 'M%' OR saldo < 1000;
""")
result_table(doc, ["nombre", "saldo"],
             [["Erick Noguera", "250000.00"], ["Maria Perez", "120000.00"]],
             "Resultado de la primera consulta (saldo >= 100000 y moneda 1):")

h(doc, "3.3 INNER JOIN entre transaccion y usuario", 2)
code_block(doc, """
SELECT t.transaction_id,
       ue.nombre AS emisor,
       ur.nombre AS receptor,
       t.importe,
       t.transaction_date
FROM transaccion t
INNER JOIN usuario ue ON ue.user_id = t.sender_user_id
INNER JOIN usuario ur ON ur.user_id = t.receiver_user_id
ORDER BY t.transaction_date;
""")
result_table(doc,
    ["transaction_id", "emisor", "receptor", "importe", "transaction_date"],
    [
        ["1", "Erick Noguera", "Maria Perez", "25000.00", "2026-08-01 10:15:00"],
        ["2", "Maria Perez", "Juan Soto", "10000.00", "2026-08-03 14:20:00"],
        ["3", "Erick Noguera", "Juan Soto", "5000.00", "2026-08-05 09:00:00"],
        ["4", "Juan Soto", "Erick Noguera", "7500.00", "2026-08-10 18:45:00"],
        ["5", "Camila Rojas", "Bruno Lima", "50.00", "2026-08-12 12:30:00"],
    ])
capture_box(doc, "Pega la captura del resultado de este JOIN ejecutado en SQLOnline.")

h(doc, "3.4 Sub-consultas: total de transacciones por usuario", 2)
code_block(doc, """
SELECT u.user_id,
       u.nombre,
       (SELECT COUNT(*)
          FROM transaccion t
         WHERE t.sender_user_id = u.user_id)      AS cantidad_enviadas,
       (SELECT COALESCE(SUM(t.importe), 0)
          FROM transaccion t
         WHERE t.sender_user_id = u.user_id)      AS total_enviado
FROM usuario u
ORDER BY total_enviado DESC;
""")
result_table(doc,
    ["user_id", "nombre", "cantidad_enviadas", "total_enviado"],
    [
        ["1", "Erick Noguera", "2", "30000.00"],
        ["2", "Maria Perez", "1", "10000.00"],
        ["3", "Juan Soto", "1", "7500.00"],
        ["4", "Camila Rojas", "1", "50.00"],
        ["5", "Bruno Lima", "0", "0.00"],
    ])

h(doc, "3.5 Funciones de agregacion con GROUP BY", 2)
code_block(doc, """
SELECT t.sender_user_id,
       COUNT(*)        AS operaciones,
       SUM(t.importe)  AS suma_importe,
       AVG(t.importe)  AS promedio_importe
FROM transaccion t
GROUP BY t.sender_user_id;
""")
result_table(doc,
    ["sender_user_id", "operaciones", "suma_importe", "promedio_importe"],
    [
        ["1", "2", "30000.00", "15000.000000"],
        ["2", "1", "10000.00", "10000.000000"],
        ["3", "1", "7500.00", "7500.000000"],
        ["4", "1", "50.00", "50.000000"],
    ])

h(doc, "3.6 Tarea Plus: vista con el top-5 de usuarios con mayor saldo", 2)
code_block(doc, """
CREATE OR REPLACE VIEW v_top5_saldo AS
SELECT u.user_id,
       u.nombre,
       u.saldo,
       m.currency_symbol
FROM usuario u
INNER JOIN moneda m ON m.currency_id = u.currency_id
ORDER BY u.saldo DESC
LIMIT 5;

SELECT * FROM v_top5_saldo;
""")
result_table(doc,
    ["user_id", "nombre", "saldo", "currency_symbol"],
    [
        ["1", "Erick Noguera", "250000.00", "$"],
        ["2", "Maria Perez", "120000.00", "$"],
        ["3", "Juan Soto", "80000.00", "$"],
        ["5", "Bruno Lima", "1500.00", "R$"],
        ["4", "Camila Rojas", "500.00", "US$"],
    ])

# =====================================================================
h(doc, "4. Leccion 3 - Manipulacion de datos (DML) y transaccionalidad (ACID)", 1)

h(doc, "4.1 Insercion de datos de prueba", 2)
code_block(doc, """
INSERT INTO moneda (currency_name, currency_symbol) VALUES
  ('Peso chileno',         '$'),
  ('Dolar estadounidense', 'US$'),
  ('Euro',                 'EUR'),
  ('Real brasileno',       'R$');

INSERT INTO usuario (nombre, correo, contrasena, saldo, currency_id) VALUES
  ('Erick Noguera', 'erick@alke.com',  'hash_pbkdf2_erick',  250000.00, 1),
  ('Maria Perez',   'maria@alke.com',  'hash_pbkdf2_maria',  120000.00, 1),
  ('Juan Soto',     'juan@alke.com',   'hash_pbkdf2_juan',    80000.00, 1),
  ('Camila Rojas',  'camila@alke.com', 'hash_pbkdf2_camila',    500.00, 2),
  ('Bruno Lima',    'bruno@alke.com',  'hash_pbkdf2_bruno',    1500.00, 4);

INSERT INTO transaccion
  (sender_user_id, receiver_user_id, currency_id, importe, transaction_date) VALUES
  (1, 2, 1, 25000.00, '2026-08-01 10:15:00'),
  (2, 3, 1, 10000.00, '2026-08-03 14:20:00'),
  (1, 3, 1,  5000.00, '2026-08-05 09:00:00'),
  (3, 1, 1,  7500.00, '2026-08-10 18:45:00'),
  (4, 5, 2,    50.00, '2026-08-12 12:30:00');
""")

h(doc, "4.2 Actualizar el saldo de un usuario y transferencia atomica", 2)
code_block(doc, """
-- Actualizacion individual del saldo tras una transaccion
UPDATE usuario SET saldo = saldo - 5000 WHERE user_id = 3;

-- Transferencia atomica: las 3 operaciones se aplican juntas o ninguna
START TRANSACTION;
  UPDATE usuario SET saldo = saldo - 15000 WHERE user_id = 1;
  UPDATE usuario SET saldo = saldo + 15000 WHERE user_id = 2;
  INSERT INTO transaccion (sender_user_id, receiver_user_id, currency_id, importe)
  VALUES (1, 2, 1, 15000.00);
COMMIT;
""")
capture_box(doc, "Pega la captura de la consola mostrando el COMMIT ejecutado con exito.")

h(doc, "4.3 Simular un error de integridad referencial y revertir", 2)
para(doc, "El receptor 999 no existe. La clave foranea fk_tx_receiver rechaza el "
          "INSERT (ERROR 1452) y el ROLLBACK deshace tambien el UPDATE previo: "
          "la wallet nunca queda con dinero descontado sin destino.")
code_block(doc, """
START TRANSACTION;
  UPDATE usuario SET saldo = saldo - 1000 WHERE user_id = 1;
  INSERT INTO transaccion (sender_user_id, receiver_user_id, currency_id, importe)
  VALUES (1, 999, 1, 1000.00);        -- falla: FK inexistente
ROLLBACK;

-- Verificacion: los saldos NO reflejan el bloque revertido
SELECT user_id, nombre, saldo FROM usuario WHERE user_id IN (1, 2, 3);
""")
result_table(doc, ["user_id", "nombre", "saldo"],
    [
        ["1", "Erick Noguera", "235000.00"],
        ["2", "Maria Perez", "135000.00"],
        ["3", "Juan Soto", "75000.00"],
    ],
    "Resultado tras la transferencia atomica confirmada y el bloque con error revertido:")

h(doc, "4.4 Propiedades ACID", 2)
info_table(doc, ["Propiedad", "Como se garantiza en Alke Wallet"],
    [
        ["Atomicidad", "START TRANSACTION ... COMMIT aplica las 2 actualizaciones de "
         "saldo y el INSERT como una sola unidad; ROLLBACK no deja cambios parciales."],
        ["Consistencia", "PK, FK, CHECK (saldo >= 0, importe > 0) y NOT NULL hacen que "
         "la base pase siempre de un estado valido a otro valido."],
        ["Aislamiento", "El motor InnoDB usa MVCC y bloqueos; una transferencia en "
         "curso no expone estados intermedios a otras transacciones concurrentes."],
        ["Durabilidad", "Tras el COMMIT los cambios quedan escritos de forma "
         "permanente (redo log de InnoDB) aunque el servidor se reinicie."],
    ])

h(doc, "4.5 Tarea Plus: cargar 50 transacciones aleatorias", 2)
code_block(doc, """
SET @n_usuarios = (SELECT COUNT(*) FROM usuario);

INSERT INTO transaccion
  (sender_user_id, receiver_user_id, currency_id, importe, transaction_date)
WITH RECURSIVE serie AS (
  SELECT 1 AS n
  UNION ALL
  SELECT n + 1 FROM serie WHERE n < 50
)
SELECT
  emisor,
  CASE WHEN receptor = emisor THEN (emisor MOD @n_usuarios) + 1 ELSE receptor END,
  1,
  ROUND(1000 + RAND(n) * 49000, 2),
  NOW() - INTERVAL FLOOR(RAND(n * 7) * 43200) MINUTE
FROM (
  SELECT n,
         FLOOR(1 + RAND(n)      * @n_usuarios) AS emisor,
         FLOOR(1 + RAND(n * 13) * @n_usuarios) AS receptor
  FROM serie
) AS base;

SELECT COUNT(*) AS total_transacciones FROM transaccion;
""")

# =====================================================================
h(doc, "5. Leccion 4 - Definicion de tablas (DDL)", 1)
para(doc, "Se crea primero moneda porque usuario y transaccion la referencian "
          "por clave foranea. Todas las tablas usan el motor InnoDB para soportar "
          "claves foraneas y transacciones.")

h(doc, "5.1 Base de datos y tabla moneda", 2)
code_block(doc, """
CREATE DATABASE IF NOT EXISTS AlkeWallet
  CHARACTER SET utf8mb4 COLLATE utf8mb4_0900_ai_ci;
USE AlkeWallet;

CREATE TABLE moneda (
  currency_id     INT           NOT NULL AUTO_INCREMENT,
  currency_name   VARCHAR(50)   NOT NULL,
  currency_symbol VARCHAR(5)    NOT NULL,
  CONSTRAINT pk_moneda        PRIMARY KEY (currency_id),
  CONSTRAINT uq_moneda_nombre UNIQUE (currency_name)
) ENGINE=InnoDB;
""")

h(doc, "5.2 Tabla usuario", 2)
code_block(doc, """
CREATE TABLE usuario (
  user_id      INT            NOT NULL AUTO_INCREMENT,
  nombre       VARCHAR(100)   NOT NULL,
  correo       VARCHAR(150)   NOT NULL,
  contrasena   VARCHAR(255)   NOT NULL,
  saldo        DECIMAL(15,2)  NOT NULL DEFAULT 0.00,
  currency_id  INT            NOT NULL,
  CONSTRAINT pk_usuario        PRIMARY KEY (user_id),
  CONSTRAINT uq_usuario_correo UNIQUE (correo),
  CONSTRAINT chk_usuario_saldo CHECK (saldo >= 0),
  CONSTRAINT fk_usuario_moneda FOREIGN KEY (currency_id)
        REFERENCES moneda (currency_id)
        ON UPDATE CASCADE ON DELETE RESTRICT
) ENGINE=InnoDB;

CREATE INDEX idx_usuario_moneda ON usuario (currency_id);
""")

h(doc, "5.3 Tabla transaccion (con indices compuestos)", 2)
code_block(doc, """
CREATE TABLE transaccion (
  transaction_id    INT            NOT NULL AUTO_INCREMENT,
  sender_user_id    INT            NOT NULL,
  receiver_user_id  INT            NOT NULL,
  currency_id       INT            NOT NULL,
  importe           DECIMAL(15,2)  NOT NULL,
  transaction_date  DATETIME       NOT NULL DEFAULT CURRENT_TIMESTAMP,
  CONSTRAINT pk_transaccion           PRIMARY KEY (transaction_id),
  CONSTRAINT chk_transaccion_importe  CHECK (importe > 0),
  CONSTRAINT chk_transaccion_distinta CHECK (sender_user_id <> receiver_user_id),
  CONSTRAINT fk_tx_sender   FOREIGN KEY (sender_user_id)
        REFERENCES usuario (user_id) ON UPDATE CASCADE ON DELETE RESTRICT,
  CONSTRAINT fk_tx_receiver FOREIGN KEY (receiver_user_id)
        REFERENCES usuario (user_id) ON UPDATE CASCADE ON DELETE RESTRICT,
  CONSTRAINT fk_tx_moneda   FOREIGN KEY (currency_id)
        REFERENCES moneda (currency_id) ON UPDATE CASCADE ON DELETE RESTRICT
) ENGINE=InnoDB;

CREATE INDEX idx_tx_sender_fecha   ON transaccion (sender_user_id, transaction_date);
CREATE INDEX idx_tx_receiver_fecha ON transaccion (receiver_user_id, transaction_date);
""")

h(doc, "5.4 Restricciones aplicadas", 2)
info_table(doc, ["Tipo", "Donde", "Proposito"],
    [
        ["PRIMARY KEY", "currency_id / user_id / transaction_id", "Identificador unico de cada fila"],
        ["FOREIGN KEY", "usuario.currency_id, transaccion.sender/receiver/currency", "Integridad referencial entre tablas"],
        ["NOT NULL", "Todos los atributos obligatorios", "Impedir datos incompletos"],
        ["UNIQUE", "usuario.correo, moneda.currency_name", "Evitar duplicados de negocio"],
        ["CHECK", "saldo >= 0, importe > 0, sender <> receiver", "Reglas de coherencia del dominio"],
        ["INDEX compuesto", "(sender_user_id, transaction_date) y receptor", "Acelerar el historial por usuario y fecha"],
    ])

h(doc, "5.5 Verificacion de la estructura", 2)
code_block(doc, """
SHOW TABLES;
DESCRIBE moneda;
DESCRIBE usuario;
DESCRIBE transaccion;
SHOW CREATE TABLE transaccion;
""")
capture_box(doc, "Pega la captura del DESCRIBE (o SHOW CREATE TABLE) de las tres tablas.")

h(doc, "5.6 Tarea Plus: anadir la fecha de creacion con ALTER TABLE", 2)
code_block(doc, """
ALTER TABLE usuario
  ADD COLUMN fecha_creacion DATETIME NOT NULL DEFAULT CURRENT_TIMESTAMP
  AFTER currency_id;
""")

# =====================================================================
h(doc, "6. Leccion 5 - El modelo entidad-relacion", 1)

h(doc, "6.1 Entidades y atributos", 2)
info_table(doc, ["Entidad", "Atributo", "Tipo", "Clave", "Nulo"],
    [
        ["moneda", "currency_id", "INT AUTO_INCREMENT", "PK", "No"],
        ["moneda", "currency_name", "VARCHAR(50)", "UNIQUE", "No"],
        ["moneda", "currency_symbol", "VARCHAR(5)", "-", "No"],
        ["usuario", "user_id", "INT AUTO_INCREMENT", "PK", "No"],
        ["usuario", "nombre", "VARCHAR(100)", "-", "No"],
        ["usuario", "correo", "VARCHAR(150)", "UNIQUE", "No"],
        ["usuario", "contrasena", "VARCHAR(255)", "-", "No"],
        ["usuario", "saldo", "DECIMAL(15,2)", "-", "No (def. 0.00)"],
        ["usuario", "currency_id", "INT", "FK -> moneda", "No"],
        ["transaccion", "transaction_id", "INT AUTO_INCREMENT", "PK", "No"],
        ["transaccion", "sender_user_id", "INT", "FK -> usuario", "No"],
        ["transaccion", "receiver_user_id", "INT", "FK -> usuario", "No"],
        ["transaccion", "currency_id", "INT", "FK -> moneda", "No"],
        ["transaccion", "importe", "DECIMAL(15,2)", "-", "No"],
        ["transaccion", "transaction_date", "DATETIME", "-", "No (def. CURRENT_TIMESTAMP)"],
    ])

h(doc, "6.2 Relaciones, cardinalidades y opcionalidad", 2)
info_table(doc, ["Relacion", "Entidad A", "A", "Entidad B", "B", "Opcionalidad"],
    [
        ["es usada por", "moneda", "1", "usuario", "N", "Todo usuario tiene 1 moneda; una moneda puede tener 0..N usuarios"],
        ["denomina", "moneda", "1", "transaccion", "N", "Toda transaccion tiene 1 moneda; una moneda 0..N transacciones"],
        ["envia (sender)", "usuario", "1", "transaccion", "N", "Toda transaccion tiene 1 emisor; un usuario 0..N enviadas"],
        ["recibe (receiver)", "usuario", "1", "transaccion", "N", "Toda transaccion tiene 1 receptor; un usuario 0..N recibidas"],
    ])
para(doc, "transaccion es una entidad asociativa: resuelve la relacion muchos a "
          "muchos 'un usuario transfiere a otro usuario' y ademas guarda atributos "
          "propios (importe, fecha). Se modela con clave primaria subrogada "
          "transaction_id. No hay entidades debiles.")

h(doc, "6.3 Diagrama entidad-relacion", 2)
img = os.path.join(AQUI, "diagrama_ER.png")
if os.path.exists(img):
    doc.add_picture(img, width=Inches(6.3))
    cap = doc.paragraphs[-1]
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
capture_box(doc, "Opcional: reemplaza o acompana con la captura del diagrama "
                 "generado en dbdiagram.io / drawSQL usando el codigo DBML de abajo.")

h(doc, "6.4 Codigo para dbdiagram.io (DBML)", 2)
code_block(doc, """
Table moneda {
  currency_id int [pk, increment]
  currency_name varchar(50) [not null, unique]
  currency_symbol varchar(5) [not null]
}

Table usuario {
  user_id int [pk, increment]
  nombre varchar(100) [not null]
  correo varchar(150) [not null, unique]
  contrasena varchar(255) [not null]
  saldo decimal(15,2) [not null, default: 0]
  currency_id int [not null]
  fecha_creacion datetime [not null]
}

Table transaccion {
  transaction_id int [pk, increment]
  sender_user_id int [not null]
  receiver_user_id int [not null]
  currency_id int [not null]
  importe decimal(15,2) [not null]
  transaction_date datetime [not null]
}

Ref: usuario.currency_id > moneda.currency_id
Ref: transaccion.sender_user_id > usuario.user_id
Ref: transaccion.receiver_user_id > usuario.user_id
Ref: transaccion.currency_id > moneda.currency_id
""")

h(doc, "6.5 Normalizacion hasta la 3FN", 2)
for titulo, texto in [
    ("Primera forma normal (1FN)",
     "Cada tabla tiene clave primaria, todos los atributos son atomicos "
     "(no hay listas ni campos compuestos) y no existen grupos repetitivos "
     "ni filas duplicadas."),
    ("Segunda forma normal (2FN)",
     "Esta en 1FN y no hay dependencias parciales: como todas las claves "
     "primarias son simples (user_id, currency_id, transaction_id), cada "
     "atributo no clave depende de la clave completa."),
    ("Tercera forma normal (3FN)",
     "Esta en 2FN y no hay dependencias transitivas. Cambio aplicado: si "
     "currency_name y currency_symbol se guardaran dentro de usuario o "
     "transaccion, existiria la dependencia transitiva "
     "user_id -> currency_id -> currency_name. Se resolvio extrayendo la "
     "entidad moneda a su propia tabla y dejando solo la clave foranea "
     "currency_id; asi currency_name y currency_symbol dependen unicamente "
     "de currency_id."),
]:
    p = doc.add_paragraph()
    p.add_run(titulo + ": ").bold = True
    p.add_run(texto)
para(doc, "Sobre saldo: es un dato derivable de las transacciones, pero el "
          "enunciado lo define como atributo de Usuario y se conserva por "
          "rendimiento; su coherencia se mantiene actualizandolo siempre dentro "
          "de transacciones ACID junto con el INSERT en transaccion.")

h(doc, "6.6 Tabla de correspondencia ER -> Relacional", 2)
info_table(doc, ["Elemento del modelo ER", "Transformacion al modelo relacional"],
    [
        ["Entidad usuario", "Tabla usuario; identificador user_id -> PRIMARY KEY"],
        ["Entidad moneda", "Tabla moneda; identificador currency_id -> PRIMARY KEY"],
        ["Entidad asociativa transaccion", "Tabla transaccion con PK subrogada transaction_id"],
        ["Atributos compuestos / multivaluados", "No existen; todos simples y monovaluados (1FN)"],
        ["Relacion moneda (1) - usuario (N)", "FK usuario.currency_id -> moneda.currency_id"],
        ["Relacion moneda (1) - transaccion (N)", "FK transaccion.currency_id -> moneda.currency_id"],
        ["Relacion usuario (1) - transaccion (N) 'envia'", "FK transaccion.sender_user_id -> usuario.user_id"],
        ["Relacion usuario (1) - transaccion (N) 'recibe'", "FK transaccion.receiver_user_id -> usuario.user_id"],
    ])

# =====================================================================
h(doc, "7. Consultas obligatorias del enunciado (resumen)", 1)

h(doc, "7.1 Nombre de la moneda elegida por un usuario especifico", 2)
code_block(doc, """
SELECT u.user_id, u.nombre, m.currency_name, m.currency_symbol
FROM usuario u
INNER JOIN moneda m ON m.currency_id = u.currency_id
WHERE u.user_id = 1;
""")
result_table(doc, ["user_id", "nombre", "currency_name", "currency_symbol"],
             [["1", "Erick Noguera", "Peso chileno", "$"]])

h(doc, "7.2 Todas las transacciones registradas", 2)
code_block(doc, """
SELECT t.transaction_id, t.sender_user_id, t.receiver_user_id,
       t.importe, m.currency_symbol, t.transaction_date
FROM transaccion t
INNER JOIN moneda m ON m.currency_id = t.currency_id
ORDER BY t.transaction_date;
""")
result_table(doc,
    ["transaction_id", "sender_user_id", "receiver_user_id", "importe", "currency_symbol", "transaction_date"],
    [
        ["1", "1", "2", "25000.00", "$", "2026-08-01 10:15:00"],
        ["2", "2", "3", "10000.00", "$", "2026-08-03 14:20:00"],
        ["3", "1", "3", "5000.00", "$", "2026-08-05 09:00:00"],
        ["4", "3", "1", "7500.00", "$", "2026-08-10 18:45:00"],
        ["5", "4", "5", "50.00", "US$", "2026-08-12 12:30:00"],
    ])

h(doc, "7.3 Todas las transacciones realizadas por un usuario especifico", 2)
code_block(doc, """
-- Enviadas por el usuario 1
SELECT t.transaction_id, t.sender_user_id, t.receiver_user_id,
       t.importe, t.transaction_date
FROM transaccion t
WHERE t.sender_user_id = 1
ORDER BY t.transaction_date;

-- Variante: enviadas O recibidas por el usuario 1
SELECT *
FROM transaccion
WHERE sender_user_id = 1 OR receiver_user_id = 1
ORDER BY transaction_date;
""")
result_table(doc,
    ["transaction_id", "sender_user_id", "receiver_user_id", "importe", "transaction_date"],
    [
        ["1", "1", "2", "25000.00", "2026-08-01 10:15:00"],
        ["3", "1", "3", "5000.00", "2026-08-05 09:00:00"],
    ],
    "Resultado de la primera consulta (solo enviadas por el usuario 1):")

h(doc, "7.4 Modificar el correo electronico de un usuario especifico", 2)
code_block(doc, """
UPDATE usuario
SET correo = 'erick.noguera@alke.com'
WHERE user_id = 1;
""")
result_table(doc, ["user_id", "nombre", "correo"],
             [["1", "Erick Noguera", "erick.noguera@alke.com"]],
             "Comprobacion con SELECT correo FROM usuario WHERE user_id = 1;")

h(doc, "7.5 Eliminar los datos de una transaccion (fila completa)", 2)
code_block(doc, """
DELETE FROM transaccion
WHERE transaction_id = 5;
""")
result_table(doc, ["transaction_id", "sender_user_id", "receiver_user_id", "importe"],
    [
        ["1", "1", "2", "25000.00"],
        ["2", "2", "3", "10000.00"],
        ["3", "1", "3", "5000.00"],
        ["4", "3", "1", "7500.00"],
    ],
    "Transacciones restantes tras el DELETE (la fila 5 desaparece):")

# =====================================================================
doc.add_page_break()
h(doc, "Anexo A - Script completo AlkeWallet.sql", 1)
para(doc, "Contenido integro del archivo database/AlkeWallet.sql, en orden de "
          "ejecucion real.")
sql_path = os.path.join(AQUI, "AlkeWallet.sql")
with open(sql_path, "r", encoding="utf-8") as f:
    code_block(doc, f.read())

# ---- Anexo B
doc.add_page_break()
h(doc, "Anexo B - Checklist de capturas de pantalla", 1)
para(doc, "El entregable pide adjuntar capturas de tu propia ejecucion en "
          "sqliteonline.com (modo MySQL 8). Reemplaza cada recuadro naranja de "
          "este documento por la imagen correspondiente:")
for item in [
    "CREATE DATABASE AlkeWallet ejecutado con exito + SHOW DATABASES; (Leccion 1).",
    "Resultado del INNER JOIN entre transaccion y usuario (Leccion 2).",
    "Consola mostrando un COMMIT exitoso de la transferencia atomica (Leccion 3).",
    "DESCRIBE o SHOW CREATE TABLE de las tres tablas (Leccion 4).",
    "Diagrama entidad-relacion completo (Leccion 5), desde dbdiagram.io o drawSQL.",
    "Resultado de al menos una de las cinco consultas obligatorias.",
]:
    doc.add_paragraph(item, style="List Number")

# --------------------------------------------------------------------- guardar
out = os.path.join(AQUI, "Alke_Wallet_Bases_de_Datos_Erick_Noguera.docx")
doc.save(out)
print("Documento generado:", out)
